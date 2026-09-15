"""Export a read-only reconciliation and usage snapshot as an Arabic review workbook."""
import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--snapshot", required=True)
    parser.add_argument("--reconciliation", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    snapshot = json.loads(Path(args.snapshot).read_text())
    result = json.loads(Path(args.reconciliation).read_text())
    book = Workbook()
    book.remove(book.active)

    def sheet(title, headers, rows):
        ws = book.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for cells in ws:
            for cell in cells:
                # Source text must never become a formula or executable link.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (
                55 if any(word in header for word in ["اسم", "ملاحظ", "ارتباط", "تفاصيل"]) else 24
            )
        return ws

    sheet("اقرأ أولاً", ["البند", "التفاصيل"], [
        ["الغرض", "تقرير مراجعة فقط؛ لم يُنفذ تحديث أو إضافة أو تعطيل أو حذف."],
        ["مصدر البيانات", "اتصال Supabase المحفوظ؛ لم يُثبت مستقلاً أنه اتصال موقع Render."],
        ["وقت القراءة UTC", snapshot["capturedAt"]],
        ["المنتجات المرجعية", result["sourceTotals"]["products"]["unique"]],
        ["المواد المرجعية", result["sourceTotals"]["warehouse"]["unique"]],
        ["المنتجات الحالية", len(snapshot["products"])],
        ["المواد الحالية", len(snapshot["warehouse"])],
        ["سياسة التنظيف", "حذف غير المستخدم فقط بعد فحص شامل؛ المرتبط يُراجع للتعطيل مع حفظ التاريخ."],
        ["الكود المتعارض المعتمد", "sk-0950 — danish sun-dried tomato دانش طماطم مجفف"],
        ["الرقم الداخلي", "يبقى ثابتاً؛ كود الملف هو SKU وليس الرقم الداخلي."],
        ["المطابقة", "تشابه الاسم لا يثبت الهوية. تطابق الكود مع اختلاف الاسم يحتاج مراجعة أيضاً."],
        ["مرشح إضافة", "يعني عدم وجود تطابق حرفي؛ قد يكون الصنف موجوداً باسم عربي أو كود مختلف."],
        ["الأرصدة والأسعار", "لا تعدَّل ولا تُصفَّر؛ ملف المرجع لا يحدد أسعار البيع أو توزيع أرصدة المستودعات."],
        ["الأقسام", "أقسام استخدام متعددة وليست بديلاً تلقائياً للفئة الأساسية؛ الأسماء غير المطابقة ليست أصنافاً جديدة."],
        ["نطاق فحص الارتباطات", f"فحص {len(snapshot['referenceChecks'])} علاقة FK فقط. غيابها لا يثبت أمان الحذف؛ قد توجد أسماء أو مراجع دون FK."],
        ["الخطوة قبل التطبيق", "اعتماد خريطة المطابقة والحالات غير المحسومة، ثم فحص أحدث الأرصدة والمراجع والمعاملات المفتوحة وأخذ نسخة استرجاع."],
    ])
    status = {
        "exact_match": "مطابقة الكود والاسم والوحدة",
        "review": "يحتاج مراجعة",
        "add_candidate": "لا تطابق حرفي — ليس قرار إضافة",
        "legacy_review": "سجل حالي يحتاج مراجعة",
    }
    issues = {
        "duplicate_current_sku": "الكود مكرر حالياً",
        "name_mismatch": "اختلاف الاسم",
        "unit_mismatch": "اختلاف الوحدة",
        "name_recode_review": "اسم مطابق بكود مختلف",
        "name_collision": "الاسم يطابق أكثر من سجل",
        "source_name_collision": "اسم مرجعي مستخدم لأكثر من كود",
        "category_unmatched": "لم يُطابق قسم حرفياً",
    }

    for namespace, title in [("products", "مطابقة المنتجات"), ("warehouse", "مطابقة المواد")]:
        rows = []
        for row in result["rows"]:
            if row["namespace"] != namespace or row["source"] is None:
                continue
            records = row["currentRecords"]
            rows.append([
                str(row["sourceCode"]), row["sourceName"], row["sourceUnit"],
                " | ".join(sorted({c["category"] for c in row["categories"]})),
                status[row["status"]],
                " | ".join(str(r["id"]) for r in records),
                " | ".join(str(r.get("sku") or "") for r in records),
                " | ".join(str(r.get("name") or "") for r in records),
                " | ".join(str(r.get("unit") or "") for r in records),
                " | ".join(issues.get(i["code"], i["code"]) for i in row["issues"]),
            ])
        sheet(title, ["الكود المرجعي", "الاسم المرجعي", "الوحدة المرجعية", "أقسام الاستخدام",
                      "نتيجة المطابقة", "الأرقام الداخلية المرشحة", "الأكواد الحالية",
                      "الأسماء الحالية", "الوحدات الحالية", "ملاحظات المراجعة"], rows)

    for namespace, table, title in [
        ("products", "products", "المنتجات الحالية"),
        ("warehouse", "warehouse_items", "المواد الحالية"),
    ]:
        rows = []
        for record in snapshot[namespace]:
            references = snapshot["references"].get(f"{table}:{record['id']}", [])
            rows.append([
                record["id"], str(record.get("sku") or ""), record["name"],
                record.get("unit"), record.get("category"), str(record.get("is_active")),
                str(record.get("current_stock", "ليس رصيداً موحداً في هذا الجدول")),
                record.get("base_price"),
                sum(r["count"] for r in references),
                " | ".join(f"{r['table']}.{r['column']}: {r['count']}" for r in references),
                "لا حذف: توجد ارتباطات" if references else "لا علاقات FK مرصودة — ليس تصريح حذف",
            ])
        sheet(title, ["الرقم الداخلي", "الكود الحالي", "الاسم الحالي", "الوحدة", "الفئة",
                      "نشط", "رصيد جدول المواد فقط", "سعر البيع الحالي إن وجد",
                      "عدد المراجع", "تفاصيل الارتباطات", "ملاحظات الحماية"], rows)

    sheet("أقسام غير مطابقة", ["اسم المصدر", "القسم", "صف المصدر"], [
        [r["name"], r["category"], r["sourceRow"]] for r in result["unmatchedCategories"]
    ])
    sheet("نطاق فحص الروابط", ["الجدول", "الحقل", "جدول الصنف"], [
        [r["table"], r["column"], r["target"]] for r in snapshot["referenceChecks"]
    ])
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(f"Created review workbook: {output}; {len(book.sheetnames)} sheets")


if __name__ == "__main__":
    main()